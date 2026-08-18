from __future__ import annotations

import asyncio
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.modules.KnowledgeBase import tasks
from src.modules.KnowledgeBase.tasks import (
    DocumentTaskContext,
    DocumentProcessingError,
    chunk_text,
    count_tokens,
    count_words,
    extract_keywords,
    extract_text,
)


@pytest.mark.parametrize("file_type", ["txt", "md"])
def test_extract_text_decodes_utf8_text_documents(file_type: str) -> None:
    content = "Leo Agent\n知识库文档"

    assert extract_text(content.encode("utf-8"), file_type) == content


def test_extract_text_strips_utf8_bom() -> None:
    content = "带 BOM 的文档"

    assert extract_text(content.encode("utf-8-sig"), "txt") == content


def test_extract_text_turns_html_into_readable_text() -> None:
    html = b"<html><body><h1>Leo</h1><p>Agent platform</p></body></html>"

    result = extract_text(html, "html")

    assert "Leo" in result
    assert "Agent platform" in result
    assert "<h1>" not in result


def test_extract_text_reads_csv_without_losing_cell_values() -> None:
    result = extract_text("name,city\nLeo,上海\n".encode(), "csv")

    assert "name" in result
    assert "city" in result
    assert "Leo" in result
    assert "上海" in result


def test_extract_text_reads_real_docx_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("Leo Agent DOCX knowledge")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "provider"
    table.cell(0, 1).text = "model"
    payload = io.BytesIO()
    document.save(payload)

    result = extract_text(payload.getvalue(), "docx")

    assert "Leo Agent DOCX knowledge" in result
    assert "provider model" in result


def test_extract_text_reads_real_pdf_text_content() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)  # noqa: SLF001
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): font_reference}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(
        b"BT /F1 12 Tf 72 720 Td (Leo Agent PDF knowledge) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(  # noqa: SLF001
        content
    )
    payload = io.BytesIO()
    writer.write(payload)

    result = extract_text(payload.getvalue(), "pdf")

    assert result == "Leo Agent PDF knowledge"


@pytest.mark.parametrize(
    ("data", "file_type"),
    [
        (b"", "txt"),
        (b"content", "exe"),
        (b"\xff\xfe", "txt"),
    ],
)
def test_extract_text_rejects_empty_unsupported_or_invalid_documents(
    data: bytes,
    file_type: str,
) -> None:
    with pytest.raises(DocumentProcessingError):
        extract_text(data, file_type)


def test_chunk_text_fixed_size_preserves_overlap_and_tail() -> None:
    chunks = chunk_text(
        "a b c d e f g h i j",
        method="fixed",
        size=4,
        overlap=1,
    )

    assert chunks == ["a b c d", "d e f g", "g h i j"]


def test_chunk_text_sentence_keeps_sentence_boundaries() -> None:
    chunks = chunk_text(
        "First sentence. Second sentence! Third sentence?",
        method="sentence",
        size=10,
        overlap=0,
    )

    assert chunks == ["First sentence. Second sentence!", "Third sentence?"]


def test_chunk_text_paragraph_keeps_paragraph_boundaries() -> None:
    chunks = chunk_text(
        "First paragraph.\n\nSecond paragraph.\n\nThird paragraph.",
        method="paragraph",
        size=12,
        overlap=0,
    )

    assert chunks == [
        "First paragraph.\n\nSecond paragraph.",
        "Third paragraph.",
    ]


def test_chunk_text_normalizes_blank_input_to_no_chunks() -> None:
    assert chunk_text(" \n\t ", method="fixed", size=10, overlap=0) == []


@pytest.mark.parametrize(
    ("method", "size", "overlap"),
    [
        ("unknown", 10, 0),
        ("fixed", 0, 0),
        ("fixed", 10, -1),
        ("fixed", 10, 10),
        ("fixed", 10, 11),
    ],
)
def test_chunk_text_rejects_unknown_method_or_invalid_window(
    method: str,
    size: int,
    overlap: int,
) -> None:
    with pytest.raises(DocumentProcessingError):
        chunk_text("some content", method=method, size=size, overlap=overlap)


def test_count_helpers_are_deterministic_for_empty_and_plain_text() -> None:
    assert count_words("") == 0
    assert count_tokens("") == 0
    assert count_words("one two\nthree") == 3
    assert count_words("知识 base") == 3
    assert count_tokens("one two") == 2


def test_extract_keywords_prefers_frequent_meaningful_terms_and_honors_limit() -> None:
    keywords = extract_keywords(
        "agent retrieval retrieval knowledge knowledge knowledge the and",
        limit=2,
    )

    assert keywords == ["knowledge", "retrieval"]


def build_task_context() -> DocumentTaskContext:
    return DocumentTaskContext(
        document_id=7,
        knowledge_base_id=3,
        file_path="kb/3/document.txt",
        file_type="txt",
        chunk_method="fixed",
        chunk_size=20,
        chunk_overlap=5,
    )


def build_async_context(value: object) -> MagicMock:
    context = MagicMock()
    context.__aenter__ = AsyncMock(return_value=value)
    context.__aexit__ = AsyncMock(return_value=None)
    return context


def build_cleanup_session(
    *,
    jobs: list[SimpleNamespace] | None = None,
    current: SimpleNamespace | None = None,
) -> MagicMock:
    session = MagicMock()
    session.begin = MagicMock(return_value=build_async_context(None))
    result = MagicMock()
    result.scalars.return_value.all.return_value = jobs or []
    session.execute = AsyncMock(return_value=result)
    session.get = AsyncMock(return_value=current)
    session.delete = AsyncMock()
    return session


@pytest.mark.asyncio
async def test_process_document_task_returns_without_io_when_claim_loses(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    claim = AsyncMock(return_value=None)
    download = MagicMock()
    complete = AsyncMock()
    failed = AsyncMock()
    monkeypatch.setattr(tasks, "_claim_document", claim)
    monkeypatch.setattr(tasks, "download_file", download)
    monkeypatch.setattr(tasks, "_complete_document", complete)
    monkeypatch.setattr(tasks, "_mark_document_failed", failed)

    await tasks.process_document_task(7)

    claim.assert_awaited_once_with(7)
    download.assert_not_called()
    complete.assert_not_awaited()
    failed.assert_not_awaited()


@pytest.mark.asyncio
async def test_process_document_task_downloads_parses_chunks_and_completes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    context = build_task_context()
    claim = AsyncMock(return_value=context)
    download = MagicMock(return_value=b"one two three four")
    extract = MagicMock(return_value="one two three four")
    chunk = MagicMock(return_value=["one two", "three four"])
    complete = AsyncMock()
    failed = AsyncMock()
    monkeypatch.setattr(tasks, "_claim_document", claim)
    monkeypatch.setattr(tasks, "download_file", download)
    monkeypatch.setattr(tasks, "extract_text", extract)
    monkeypatch.setattr(tasks, "chunk_text", chunk)
    monkeypatch.setattr(tasks, "_complete_document", complete)
    monkeypatch.setattr(tasks, "_mark_document_failed", failed)

    await tasks.process_document_task(context.document_id)

    claim.assert_awaited_once_with(context.document_id)
    download.assert_called_once_with(context.file_path)
    extract.assert_called_once_with(b"one two three four", context.file_type)
    chunk.assert_called_once_with(
        "one two three four",
        context.chunk_method,
        context.chunk_size,
        context.chunk_overlap,
    )
    complete.assert_awaited_once()
    completed_context, completed_text, processed_segments = (
        complete.await_args.args
    )
    assert completed_context is context
    assert completed_text == "one two three four"
    assert [segment.content for segment in processed_segments] == [
        "one two",
        "three four",
    ]
    assert all(segment.word_count == 2 for segment in processed_segments)
    assert all(segment.token_count > 0 for segment in processed_segments)
    failed.assert_not_awaited()


@pytest.mark.asyncio
async def test_process_document_task_marks_document_failed_and_swallows_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    context = build_task_context()
    claim = AsyncMock(return_value=context)
    download = MagicMock(side_effect=RuntimeError("MinIO unavailable"))
    complete = AsyncMock()
    failed = AsyncMock()
    monkeypatch.setattr(tasks, "_claim_document", claim)
    monkeypatch.setattr(tasks, "download_file", download)
    monkeypatch.setattr(tasks, "_complete_document", complete)
    monkeypatch.setattr(tasks, "_mark_document_failed", failed)

    await tasks.process_document_task(context.document_id)

    complete.assert_not_awaited()
    failed.assert_awaited_once()
    failed_document_id, error = failed.await_args.args
    assert failed_document_id == context.document_id
    assert isinstance(error, str)
    assert "RuntimeError" in error
    assert "MinIO unavailable" not in error


@pytest.mark.asyncio
async def test_process_document_task_records_cancellation_then_reraises(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    context = build_task_context()
    claim = AsyncMock(return_value=context)
    download = MagicMock(side_effect=asyncio.CancelledError)
    failed = AsyncMock()
    monkeypatch.setattr(tasks, "_claim_document", claim)
    monkeypatch.setattr(tasks, "download_file", download)
    monkeypatch.setattr(tasks, "_mark_document_failed", failed)

    with pytest.raises(asyncio.CancelledError):
        await tasks.process_document_task(context.document_id)

    failed.assert_awaited_once_with(
        context.document_id,
        "文档处理任务已取消",
    )


@pytest.mark.asyncio
async def test_process_storage_cleanup_jobs_deletes_object_and_job_record(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    job = SimpleNamespace(
        id=11,
        object_name="kb/3/obsolete.txt",
        attempts=0,
        last_error=None,
    )
    list_session = build_cleanup_session(jobs=[job])
    delete_session = build_cleanup_session(current=job)
    session_factory = MagicMock(
        side_effect=[
            build_async_context(list_session),
            build_async_context(delete_session),
        ]
    )
    minio_delete = MagicMock()
    monkeypatch.setattr(tasks, "AsyncSessionLocal", session_factory)
    monkeypatch.setattr(tasks, "delete_file", minio_delete)

    cleaned = await tasks.process_storage_cleanup_jobs([job.id])

    assert cleaned == 1
    minio_delete.assert_called_once_with(job.object_name)
    delete_session.get.assert_awaited_once_with(
        tasks.StorageCleanupJob,
        job.id,
        with_for_update=True,
    )
    delete_session.delete.assert_awaited_once_with(job)
    assert session_factory.call_count == 2


@pytest.mark.asyncio
async def test_process_storage_cleanup_jobs_retains_failed_job_with_safe_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    selected_job = SimpleNamespace(
        id=12,
        object_name="kb/3/retry.txt",
        attempts=2,
        last_error=None,
    )
    current_job = SimpleNamespace(
        id=selected_job.id,
        object_name=selected_job.object_name,
        attempts=2,
        last_error=None,
    )
    list_session = build_cleanup_session(jobs=[selected_job])
    update_session = build_cleanup_session(current=current_job)
    session_factory = MagicMock(
        side_effect=[
            build_async_context(list_session),
            build_async_context(update_session),
        ]
    )
    raw_error = "MinIO secret credential leaked from transport"
    minio_delete = MagicMock(side_effect=RuntimeError(raw_error))
    monkeypatch.setattr(tasks, "AsyncSessionLocal", session_factory)
    monkeypatch.setattr(tasks, "delete_file", minio_delete)

    cleaned = await tasks.process_storage_cleanup_jobs([selected_job.id])

    assert cleaned == 0
    minio_delete.assert_called_once_with(selected_job.object_name)
    update_session.get.assert_awaited_once_with(
        tasks.StorageCleanupJob,
        selected_job.id,
        with_for_update=True,
    )
    assert current_job.attempts == 3
    assert current_job.last_error == "MinIO 清理失败（RuntimeError）"
    assert raw_error not in current_job.last_error
    update_session.delete.assert_not_awaited()
    assert session_factory.call_count == 2


@pytest.mark.asyncio
async def test_process_storage_cleanup_jobs_empty_ids_is_noop(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    session_factory = MagicMock()
    minio_delete = MagicMock()
    monkeypatch.setattr(tasks, "AsyncSessionLocal", session_factory)
    monkeypatch.setattr(tasks, "delete_file", minio_delete)

    cleaned = await tasks.process_storage_cleanup_jobs([])

    assert cleaned == 0
    session_factory.assert_not_called()
    minio_delete.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("document_count", "statuses", "expected_status"),
    [
        (0, [], "empty"),
        (2, [("completed", 2)], "ready"),
        (2, [("completed", 1), ("failed", 1)], "ready"),
        (2, [("failed", 1), ("pending", 1)], "indexing"),
        (2, [("completed", 1), ("processing", 1)], "ready"),
    ],
)
async def test_recalculate_knowledge_base_uses_aggregate_status_precedence(
    document_count: int,
    statuses: list[tuple[str, int]],
    expected_status: str,
) -> None:
    db = MagicMock()
    db.scalar = AsyncMock(side_effect=[document_count, 7])
    status_result = MagicMock()
    status_result.all.return_value = statuses
    db.execute = AsyncMock(return_value=status_result)
    knowledge_base = SimpleNamespace(
        id=3,
        document_count=-1,
        segment_count=-1,
        status="unknown",
    )

    await tasks._recalculate_knowledge_base(db, knowledge_base)

    assert knowledge_base.document_count == document_count
    assert knowledge_base.segment_count == 7
    assert knowledge_base.status == expected_status
    assert db.scalar.await_count == 2
    db.execute.assert_awaited_once()
