"""In-process document ingestion tasks for the knowledge-base module.

FastAPI ``BackgroundTasks`` is intentionally used as a lightweight first
implementation.  The task receives only a document id and owns every database
transaction it needs, so it never depends on a request-scoped session.
"""

from __future__ import annotations

import asyncio
import csv
import io
import re
import zipfile
from bisect import bisect_right
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser

from loguru import logger
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.core.config import get_settings
from src.infra.database import AsyncSessionLocal
from src.infra.minio_client import delete_file, download_file
from src.modules.KnowledgeBase.model import (
    Document,
    KnowledgeBase,
    Segment,
    StorageCleanupJob,
)

SUPPORTED_DOCUMENT_TYPES = frozenset({"txt", "md", "csv", "html", "docx", "pdf"})
MAX_EXTRACTED_CHARACTERS = 10_000_000
MAX_PDF_PAGES = 500
MAX_DOCX_ENTRIES = 10_000
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_SEGMENTS = get_settings().KNOWLEDGE_MAX_SEGMENTS
MAX_ERROR_MESSAGE_LENGTH = 2_000

_TOKEN_PATTERN = re.compile(
    r"[\u3400-\u4dbf\u4e00-\u9fff]|[A-Za-z0-9_]{1,4}|[^\s]"
)
_WORD_PATTERN = re.compile(
    r"[\u3400-\u4dbf\u4e00-\u9fff]|[A-Za-z0-9]+(?:[-_'][A-Za-z0-9]+)*"
)
_KEYWORD_PATTERN = re.compile(
    r"[\u3400-\u4dbf\u4e00-\u9fff]{2,8}|[A-Za-z][A-Za-z0-9_-]{1,31}"
)
_SENTENCE_BOUNDARY_PATTERN = re.compile(
    r"[。！？!?；;]+(?:[\"'”’）】》」』]*)|\.(?=\s|$)|\n+"
)
_PARAGRAPH_BOUNDARY_PATTERN = re.compile(r"\n\s*\n+")

_STOP_WORDS = frozenset(
    {
        "and",
        "are",
        "for",
        "from",
        "into",
        "that",
        "the",
        "this",
        "was",
        "were",
        "with",
        "一个",
        "以及",
        "但是",
        "可以",
        "如果",
        "已经",
        "我们",
        "所有",
        "这个",
        "这些",
        "进行",
        "通过",
    }
)


class DocumentProcessingError(ValueError):
    """A safe, user-displayable document-processing failure."""


@dataclass(frozen=True, slots=True)
class DocumentTaskContext:
    document_id: int
    knowledge_base_id: int
    file_path: str
    file_type: str
    chunk_method: str
    chunk_size: int
    chunk_overlap: int


@dataclass(frozen=True, slots=True)
class ProcessedSegment:
    content: str
    word_count: int
    token_count: int
    keywords: list[str]


class _TextHTMLParser(HTMLParser):
    """Extract visible text without introducing an HTML parsing dependency."""

    _BLOCK_TAGS = frozenset(
        {
            "address",
            "article",
            "aside",
            "blockquote",
            "br",
            "dd",
            "div",
            "dl",
            "dt",
            "figcaption",
            "figure",
            "footer",
            "h1",
            "h2",
            "h3",
            "h4",
            "h5",
            "h6",
            "header",
            "hr",
            "li",
            "main",
            "nav",
            "ol",
            "p",
            "pre",
            "section",
            "table",
            "td",
            "th",
            "tr",
            "ul",
        }
    )
    _IGNORED_TAGS = frozenset({"script", "style", "noscript", "template"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        del attrs
        normalized_tag = tag.lower()
        if normalized_tag in self._IGNORED_TAGS:
            self._ignored_depth += 1
        elif not self._ignored_depth and normalized_tag in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        normalized_tag = tag.lower()
        if normalized_tag in self._IGNORED_TAGS:
            self._ignored_depth = max(0, self._ignored_depth - 1)
        elif not self._ignored_depth and normalized_tag in self._BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self._parts.append(data)

    def text(self) -> str:
        return "".join(self._parts)


def count_words(text: str) -> int:
    """Return a language-neutral approximate word count.

    Each CJK character counts as one word, while contiguous Latin letters and
    digits count as one word.  The value is deterministic but is not intended
    to be a linguistic tokenizer result.
    """

    return sum(1 for _ in _WORD_PATTERN.finditer(text))


def count_tokens(text: str) -> int:
    """Estimate model tokens without coupling ingestion to one tokenizer."""

    return sum(1 for _ in _TOKEN_PATTERN.finditer(text))


def extract_keywords(text: str, limit: int = 8) -> list[str]:
    """Extract deterministic high-frequency terms for display and retrieval."""

    if limit <= 0:
        return []
    normalized_terms = (
        term.casefold()
        for term in _KEYWORD_PATTERN.findall(text)
        if term.casefold() not in _STOP_WORDS
    )
    counts = Counter(normalized_terms)
    return [
        term
        for term, _ in sorted(
            counts.items(),
            key=lambda item: (-item[1], -len(item[0]), item[0]),
        )[:limit]
    ]


def extract_text(data: bytes, file_type: str) -> str:
    """Extract normalized text from a supported document payload."""

    normalized_type = file_type.lower().lstrip(".")
    if normalized_type not in SUPPORTED_DOCUMENT_TYPES:
        raise DocumentProcessingError(f"不支持的文件类型: {file_type}")
    if not data:
        raise DocumentProcessingError("文档内容为空")

    try:
        if normalized_type in {"txt", "md"}:
            text = _decode_text(data)
        elif normalized_type == "csv":
            text = _extract_csv(data)
        elif normalized_type == "html":
            text = _extract_html(data)
        elif normalized_type == "docx":
            text = _extract_docx(data)
        else:
            text = _extract_pdf(data)
    except DocumentProcessingError:
        raise
    except Exception as exc:
        raise DocumentProcessingError(
            f"{normalized_type.upper()} 文档解析失败"
        ) from exc

    normalized = _normalize_text(text)
    if not normalized:
        hint = "，扫描版 PDF 需要 OCR" if normalized_type == "pdf" else ""
        raise DocumentProcessingError(f"文档中没有可提取的文本{hint}")
    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentProcessingError("文档解析后的文本超过处理上限")
    return normalized


def chunk_text(
    text: str,
    method: str,
    size: int,
    overlap: int,
) -> list[str]:
    """Split text into bounded, optionally overlapping approximate-token chunks."""

    normalized_method = method.lower()
    if normalized_method not in {"fixed", "sentence", "paragraph"}:
        raise DocumentProcessingError(f"不支持的分段方式: {method}")
    if size <= 0:
        raise DocumentProcessingError("分段大小必须大于 0")
    if overlap < 0 or overlap >= size:
        raise DocumentProcessingError("分段重叠必须大于等于 0 且小于分段大小")

    normalized_text = _normalize_text(text)
    token_matches = list(_TOKEN_PATTERN.finditer(normalized_text))
    if not token_matches:
        return []

    preferred_boundaries: list[int] = []
    if normalized_method != "fixed":
        boundary_pattern = (
            _SENTENCE_BOUNDARY_PATTERN
            if normalized_method == "sentence"
            else _PARAGRAPH_BOUNDARY_PATTERN
        )
        token_ends = [match.end() for match in token_matches]
        preferred_boundaries = sorted(
            {
                bisect_right(token_ends, boundary.end())
                for boundary in boundary_pattern.finditer(normalized_text)
            }
            | {len(token_matches)}
        )

    chunks: list[str] = []
    start = 0
    previous_end = -1
    token_total = len(token_matches)
    while start < token_total:
        maximum_end = min(start + size, token_total)
        end = maximum_end
        if preferred_boundaries and maximum_end < token_total:
            boundary_index = bisect_right(preferred_boundaries, maximum_end) - 1
            if boundary_index >= 0:
                preferred_end = preferred_boundaries[boundary_index]
                if preferred_end > start and preferred_end != previous_end:
                    end = preferred_end

        char_start = token_matches[start].start()
        char_end = token_matches[end - 1].end()
        chunk = normalized_text[char_start:char_end].strip()
        if chunk:
            chunks.append(chunk)
            if len(chunks) > MAX_SEGMENTS:
                raise DocumentProcessingError("文档分段数量超过处理上限")
        if end >= token_total:
            break
        previous_end = end
        start = max(start + 1, end - overlap)

    return chunks


async def process_document_task(document_id: int) -> None:
    """Process one uploaded document using task-owned database sessions."""

    context: DocumentTaskContext | None = None
    try:
        context = await _claim_document(document_id)
        if context is None:
            return

        file_bytes = await asyncio.to_thread(download_file, context.file_path)
        text = await asyncio.to_thread(extract_text, file_bytes, context.file_type)
        chunks = await asyncio.to_thread(
            chunk_text,
            text,
            context.chunk_method,
            context.chunk_size,
            context.chunk_overlap,
        )
        if not chunks:
            raise DocumentProcessingError("文档未生成有效分段")
        processed_segments = await asyncio.to_thread(_prepare_segments, chunks)
        await _complete_document(context, text, processed_segments)
    except asyncio.CancelledError:
        if context is not None:
            await _mark_document_failed(document_id, "文档处理任务已取消")
        raise
    except Exception as exc:
        logger.exception("Document processing failed: document_id={}", document_id)
        await _mark_document_failed(document_id, _safe_error_message(exc))


async def process_storage_cleanup_jobs(
    job_ids: list[int] | None = None,
    *,
    limit: int = 100,
) -> int:
    """Best-effort processing for durable MinIO cleanup records.

    Successful work is deleted. Failed work remains in MySQL with a safe error
    and can be retried on the next application start or explicit call.
    """

    if limit <= 0:
        raise ValueError("limit 必须大于 0")
    if job_ids == []:
        return 0

    effective_limit = max(limit, len(job_ids)) if job_ids is not None else limit
    async with AsyncSessionLocal() as db:
        stmt = (
            select(StorageCleanupJob)
            .order_by(StorageCleanupJob.id)
            .limit(effective_limit)
        )
        if job_ids is not None:
            stmt = stmt.where(StorageCleanupJob.id.in_(list(dict.fromkeys(job_ids))))
        jobs = list((await db.execute(stmt)).scalars().all())

    cleaned = 0
    for job in jobs:
        try:
            await asyncio.to_thread(delete_file, job.object_name)
        except Exception as exc:
            async with AsyncSessionLocal() as db:
                async with db.begin():
                    current = await db.get(
                        StorageCleanupJob,
                        job.id,
                        with_for_update=True,
                    )
                    if current is not None:
                        current.attempts += 1
                        current.last_error = (
                            f"MinIO 清理失败（{type(exc).__name__}）"
                        )[:500]
            logger.warning(
                "MinIO cleanup deferred: job_id={}, object_name={}",
                job.id,
                job.object_name,
            )
            continue

        async with AsyncSessionLocal() as db:
            async with db.begin():
                current = await db.get(
                    StorageCleanupJob,
                    job.id,
                    with_for_update=True,
                )
                if current is not None:
                    await db.delete(current)
        cleaned += 1
    return cleaned


async def _claim_document(document_id: int) -> DocumentTaskContext | None:
    """Atomically claim a pending or failed document for processing."""

    knowledge_base_id = await _lookup_document_knowledge_base_id(document_id)
    if knowledge_base_id is None:
        logger.warning(
            "Document task skipped because the document does not exist: {}",
            document_id,
        )
        return None
    async with AsyncSessionLocal() as db:
        async with db.begin():
            kb = await _get_knowledge_base_for_update(db, knowledge_base_id)
            if kb is None:
                logger.warning(
                    "Document task skipped because its knowledge base does not exist: {}",
                    document_id,
                )
                return None
            doc = await _get_document_for_update(db, document_id)
            if doc is None or doc.knowledge_base_id != knowledge_base_id:
                logger.warning(
                    "Document task skipped because the document does not exist: {}",
                    document_id,
                )
                return None
            if doc.status in {"processing", "completed"}:
                logger.info(
                    "Document task skipped because status is {}: {}",
                    doc.status,
                    document_id,
                )
                return None
            if doc.status not in {"pending", "failed"}:
                logger.warning(
                    "Document task skipped because status is invalid: id={}, status={}",
                    document_id,
                    doc.status,
                )
                return None
            if not doc.file_path:
                raise DocumentProcessingError("文档缺少 MinIO 存储路径")

            doc.status = "processing"
            doc.error_message = None
            doc.processed_at = None
            await db.flush()
            await _recalculate_knowledge_base(db, kb)
            return DocumentTaskContext(
                document_id=doc.id,
                knowledge_base_id=doc.knowledge_base_id,
                file_path=doc.file_path,
                file_type=doc.file_type,
                chunk_method=kb.chunk_method,
                chunk_size=kb.chunk_size,
                chunk_overlap=kb.chunk_overlap,
            )


async def _complete_document(
    context: DocumentTaskContext,
    text: str,
    segments: list[ProcessedSegment],
) -> None:
    """Replace a document's segments and finalize counters in one transaction."""

    async with AsyncSessionLocal() as db:
        async with db.begin():
            kb = await _get_knowledge_base_for_update(
                db,
                context.knowledge_base_id,
            )
            if kb is None:
                return
            doc = await _get_document_for_update(db, context.document_id)
            if (
                doc is None
                or doc.knowledge_base_id != context.knowledge_base_id
            ):
                logger.info(
                    "Document disappeared before processing completed: {}",
                    context.document_id,
                )
                return
            if doc.status != "processing":
                logger.info(
                    "Document finalization skipped because status changed: id={}, status={}",
                    context.document_id,
                    doc.status,
                )
                return

            await db.execute(
                delete(Segment).where(Segment.document_id == context.document_id)
            )
            db.add_all(
                [
                    Segment(
                        knowledge_base_id=context.knowledge_base_id,
                        document_id=context.document_id,
                        position=position,
                        content=segment.content,
                        word_count=segment.word_count,
                        token_count=segment.token_count,
                        keywords=segment.keywords,
                        hit_count=0,
                    )
                    for position, segment in enumerate(segments)
                ]
            )
            doc.status = "completed"
            doc.segment_count = len(segments)
            doc.word_count = count_words(text)
            doc.error_message = None
            doc.processed_at = datetime.now()
            await db.flush()
            await _recalculate_knowledge_base(db, kb)


async def _mark_document_failed(document_id: int, error: str) -> None:
    """Persist a failed state without relying on the failed task transaction."""

    try:
        knowledge_base_id = await _lookup_document_knowledge_base_id(document_id)
        if knowledge_base_id is None:
            return
        async with AsyncSessionLocal() as db:
            async with db.begin():
                kb = await _get_knowledge_base_for_update(
                    db,
                    knowledge_base_id,
                )
                if kb is None:
                    return
                doc = await _get_document_for_update(db, document_id)
                if (
                    doc is None
                    or doc.knowledge_base_id != knowledge_base_id
                    or doc.status == "completed"
                ):
                    return
                doc.status = "failed"
                doc.segment_count = 0
                doc.word_count = 0
                doc.error_message = error[:MAX_ERROR_MESSAGE_LENGTH]
                doc.processed_at = None
                await db.execute(
                    delete(Segment).where(Segment.document_id == document_id)
                )
                await db.flush()
                await _recalculate_knowledge_base(db, kb)
    except Exception:
        logger.exception(
            "Failed to persist document task failure: document_id={}",
            document_id,
        )


async def _get_document_for_update(
    db: AsyncSession,
    document_id: int,
) -> Document | None:
    result = await db.execute(
        select(Document).where(Document.id == document_id).with_for_update()
    )
    return result.scalar_one_or_none()


async def _get_document_knowledge_base_id(
    db: AsyncSession,
    document_id: int,
) -> int | None:
    return await db.scalar(
        select(Document.knowledge_base_id).where(Document.id == document_id)
    )


async def _lookup_document_knowledge_base_id(document_id: int) -> int | None:
    """Resolve the parent in a short read transaction before lock ordering."""

    async with AsyncSessionLocal() as db:
        return await _get_document_knowledge_base_id(db, document_id)


async def _get_knowledge_base_for_update(
    db: AsyncSession,
    knowledge_base_id: int,
) -> KnowledgeBase | None:
    result = await db.execute(
        select(KnowledgeBase)
        .where(KnowledgeBase.id == knowledge_base_id)
        .with_for_update()
    )
    return result.scalar_one_or_none()


async def _recalculate_knowledge_base(
    db: AsyncSession,
    kb: KnowledgeBase,
) -> None:
    """Recompute aggregate values to remain correct under concurrent tasks."""

    document_count = await db.scalar(
        select(func.count(Document.id)).where(
            Document.knowledge_base_id == kb.id
        )
    )
    segment_count = await db.scalar(
        select(func.count(Segment.id)).where(Segment.knowledge_base_id == kb.id)
    )
    status_rows = await db.execute(
        select(Document.status, func.count(Document.id))
        .where(Document.knowledge_base_id == kb.id)
        .group_by(Document.status)
    )
    statuses = {status: count for status, count in status_rows.all()}

    kb.document_count = int(document_count or 0)
    kb.segment_count = int(segment_count or 0)
    if kb.document_count == 0:
        kb.status = "empty"
    elif statuses.get("completed", 0):
        kb.status = "ready"
    elif statuses.get("pending", 0) or statuses.get("processing", 0):
        kb.status = "indexing"
    elif statuses.get("failed", 0):
        kb.status = "error"
    else:
        kb.status = "empty"


def _prepare_segments(chunks: list[str]) -> list[ProcessedSegment]:
    return [
        ProcessedSegment(
            content=chunk,
            word_count=count_words(chunk),
            token_count=count_tokens(chunk),
            keywords=extract_keywords(chunk),
        )
        for chunk in chunks
    ]


def _decode_text(data: bytes) -> str:
    encodings = ["utf-8-sig"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.append("utf-16")
    encodings.append("gb18030")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentProcessingError("无法识别文本文档编码")


def _extract_csv(data: bytes) -> str:
    rows = csv.reader(io.StringIO(_decode_text(data), newline=None))
    normalized_rows = [
        "\t".join(cell.strip() for cell in row if cell.strip())
        for row in rows
    ]
    return "\n".join(row for row in normalized_rows if row)


def _extract_html(data: bytes) -> str:
    parser = _TextHTMLParser()
    parser.feed(_decode_text(data))
    parser.close()
    return parser.text()


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise DocumentProcessingError("DOCX 解析依赖 python-docx 未安装") from exc

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            uncompressed_size = sum(entry.file_size for entry in entries)
            if len(entries) > MAX_DOCX_ENTRIES:
                raise DocumentProcessingError("DOCX 文件条目数量超过处理上限")
            if uncompressed_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentProcessingError("DOCX 解压后大小超过处理上限")
            if "word/document.xml" not in archive.namelist():
                raise DocumentProcessingError("DOCX 文档结构无效")
    except zipfile.BadZipFile as exc:
        raise DocumentProcessingError("DOCX 文档结构无效") from exc

    document = DocxDocument(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentProcessingError("PDF 解析依赖 pypdf 未安装") from exc

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:
            raise DocumentProcessingError("加密 PDF 无法解析") from exc
        if not unlocked:
            raise DocumentProcessingError("加密 PDF 无法解析")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentProcessingError("PDF 页数超过处理上限")
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    normalized = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", normalized).strip()


def _safe_error_message(exc: Exception) -> str:
    if isinstance(exc, DocumentProcessingError):
        return str(exc)[:MAX_ERROR_MESSAGE_LENGTH]
    return f"文档处理失败（{type(exc).__name__}）"
